#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
美术文档 → PPT 生成器（原生图表 + 真实色板，可在 PowerPoint / WPS 中直接编辑）

用法：
    python build_art_ppt.py <data.json> [output.pptx]

依赖：
    pip install python-pptx>=0.6.21

设计原则：
  1. 只吃 JSON，不自己造内容 —— 数据缺失的板块整页跳过，并记入「下一步与待定」页
  2. 图表用 python-pptx 原生 chart，不截图 —— 生成后可在 PPT 里改数据改样式
  3. 色彩板页用**真实色号填单元格底色**，并自动选对比文字色，保证可读
  4. 风格锚点页必含「禁忌」，不可省略

公共底座见同目录 ppt_kit.py（**勿单独修改副本**，由 sync_ppt_kit.py 统一同步）。
"""

from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE

import ppt_kit as K
from ppt_kit import (DataError, SH, SW, C_MUTED, C_TEXT, C_WHITE, add_chart, add_table,
                     caption, check_dict_list, check_meta, footer, new_page, new_presentation,
                     para, run_cli)

# 本脚本的排版参数（青碧主色，与另两个专家区分）
K.configure(
    accent=RGBColor(0x1F, 0x6F, 0x6B),
    title_color=RGBColor(0x1B, 0x32, 0x30),
    header_bg=RGBColor(0x1F, 0x6F, 0x6B),
    band_bg=RGBColor(0xED, 0xF5, 0xF4),
    title_y=0.4,
    title_size=25,
    bar_y=1.2,
    table_y=1.6,
    chart_y=1.6,
    chart_bottom_margin=1.4,
    caption_y=SH - 1.28,
    footer_y=SH - 0.58,
    footer_h=0.32,
)
C_ACCENT = K.C_ACCENT
C_TITLE = K.C_TITLE


def _hex_to_rgb(s):
    """'#RRGGBB' / 'RRGGBB' → RGBColor；非法返回 None。"""
    if not isinstance(s, str):
        return None
    t = s.strip().lstrip('#')
    if len(t) != 6:
        return None
    try:
        return RGBColor.from_string(t.upper())
    except Exception:
        return None


def _contrast_text(s):
    """按色块明度选白字或深字，避免色号文字看不清。"""
    c = _hex_to_rgb(s)
    if c is None:
        return None
    lum = 0.299 * c[0] + 0.587 * c[1] + 0.114 * c[2]
    return C_WHITE if lum < 140 else C_TEXT


def _chip_fill(txt, row, col):
    """只在色彩板的「色号」列填真实底色。"""
    return _hex_to_rgb(txt) if col == 1 else None


def _cap(y, row_h):
    """按可用纵向空间算表格能放多少行数据（含表头一行）。

    行数不设上限时，长清单会把表格撑到画布外（实测 30 条问题时表高 10.54 英寸，
    画布只有 7.5 英寸）—— 产物仍然"生成成功"，是最难发现的一类问题。
    """
    avail = K.CONFIG['caption_y'] - y
    return max(1, int(avail / row_h) - 1)


def _footer_text(data):
    meta = data.get('meta', {})
    n_next = len(data.get('next') or [])
    return '版本 %s ｜ %s ｜ %s / %s ｜ 待定事项 %d 项' % (
        meta.get('version', 'v1'), meta.get('date', '未标注'),
        meta.get('resolution', '未标注'), meta.get('platform', '未标注'), n_next)


DOC_KINDS = ('风格提案', '评审报告', '资产清单')   # 与 deliverables.md 规定的三类文档一致


def _default_out_name(data):
    """默认输出名与文档规定的三类命名对齐（由 meta.doc_type 决定，缺省按风格提案）。"""
    meta = data.get('meta') or {}
    kind = str(meta.get('doc_type') or '').strip()
    if kind not in DOC_KINDS:
        kind = DOC_KINDS[0]
    return '%s-美术%s.pptx' % (meta.get('project', 'art'), kind)


def coerce(data):
    """校验并归一化数据。缺失板块整页跳过；字段值不合法则逐条报错、拒绝出图。"""
    p = []
    if not isinstance(data, dict):
        raise DataError(['顶层必须是 JSON 对象'])
    check_meta(data, p)

    for key in ('anchors', 'palette', 'luminance', 'assets', 'batches', 'gaps'):
        check_dict_list(data, key, p)

    # next 是「字符串数组」，不是对象数组
    nxt = data.get('next')
    if nxt is not None:
        if not isinstance(nxt, (list, tuple)):
            p.append('next：期望数组，实际是 %s' % type(nxt).__name__)
        else:
            for i, x in enumerate(nxt):
                if not isinstance(x, str):
                    p.append('next[%d]：期望字符串，实际是 %s' % (i, type(x).__name__))

    for key in ('palette', 'assets', 'batches'):
        rows = data.get(key)
        if not isinstance(rows, (list, tuple)):
            continue
        w = {'palette': '色号', 'assets': '数量', 'batches': '数量'}[key]
        for i, x in enumerate(rows):
            if not isinstance(x, dict):
                continue
            v = x.get('count' if key != 'palette' else 'hex')
            if key == 'palette':
                if v is not None and _hex_to_rgb(v) is None:
                    p.append('palette[%d].hex：期望 #RRGGBB，实际是 %r' % (i, v))
            else:
                if v is None:
                    p.append('%s[%d].count：缺少%s' % (key, i, w))
                elif isinstance(v, bool) or not isinstance(v, (int, float)):
                    try:
                        x['count'] = float(str(v).strip())
                    except ValueError:
                        p.append('%s[%d].count：期望数值，实际是 %r' % (key, i, v))

    rv = data.get('review')
    if rv is not None and not isinstance(rv, dict):
        p.append('review：期望对象')
    elif isinstance(rv, dict):
        check_dict_list(rv, 'issues', p, 'review.issues')

    meta = data.get('meta')
    if isinstance(meta, dict) and meta.get('accent'):
        c = _hex_to_rgb(meta['accent'])
        if c is None:
            p.append('meta.accent：期望 #RRGGBB，实际是 %r' % meta['accent'])
    if isinstance(meta, dict) and meta.get('doc_type') not in (None, *DOC_KINDS):
        p.append('meta.doc_type：只支持 %s，实际是 %r' % ('/'.join(DOC_KINDS), meta['doc_type']))
    return data, p


def build(data):
    data, problems = coerce(data)
    if problems:
        raise DataError(problems)
    # 主色在此处设置（不放在 coerce 里）：校验函数不该有副作用，
    # 否则同一进程多次调用时，上一次的 accent 会残留到下一次
    _acc = _hex_to_rgb((data.get('meta') or {}).get('accent') or '')
    K.configure(accent=_acc or RGBColor(0x1F, 0x6F, 0x6B))
    prs, blank = new_presentation()
    meta = data.get('meta', {})
    skipped = []
    foot = _footer_text(data)

    # 1 封面
    s = prs.slides.add_slide(blank)
    tf = K.txbox(s, 1.1, 2.15, SW - 2.2, 3.2)
    para(tf, '%s 美术风格提案' % meta.get('project', '未命名项目'), size=40, bold=True,
         color=C_TITLE, first=True, space_after=10)
    para(tf, '%s ｜ %s ｜ %s' % (meta.get('resolution', '—'), meta.get('platform', '—'),
                                 meta.get('version', 'v1')), size=14, color=C_MUTED, space_after=18)
    para(tf, meta.get('tagline', '（未提供一句话风格定位）'), size=17, color=C_ACCENT, space_after=8)
    K.accent_bar(s, 1.1, 1.95, 1.6, 0.07)
    footer(s, foot)

    # 2 风格锚点
    an = data.get('anchors') or []
    if an:
        s = new_page(prs, blank, '风格锚点', '生成时整段逐字复用；「禁忌」不可省略')
        add_table(s, ['项目', '内容'], [[a.get('item', ''), a.get('value', '')] for a in an],
                  col_w=[1, 3], font_size=12.5, row_h=0.42, max_rows=_cap(1.6, 0.42))
        caption(s, data.get('anchors_note', '风格锚点是唯一风格来源，后续所有产出与评审都以它为依据。'))
        footer(s, foot)
    else:
        skipped.append('风格锚点')

    # 3 色彩板（真实色块）
    pal = data.get('palette') or []
    if pal:
        s = new_page(prs, blank, '色彩板', '色号直接写入规格文档，不允许凭感觉调')
        rows = [[p.get('name', ''), p.get('hex', ''), p.get('usage', ''), p.get('luminance', '')]
                for p in pal]
        add_table(s, ['色名', '色号', '用途', '明度档'], rows,
                  col_w=[1, 1, 2, 0.8], font_size=11.5, row_h=0.38,
                  cell_fill=_chip_fill, cell_color=lambda txt, c: _contrast_text(txt) if c == 1 else None,
                  max_rows=_cap(1.6, 0.38))
        caption(s, data.get('palette_note', '主体与背景的明度必须错开至少 2 档，否则缩略图下会糊。'))
        footer(s, foot)
    else:
        skipped.append('色彩板')

    # 4 明度阶梯
    lum = data.get('luminance') or []
    if lum:
        s = new_page(prs, blank, '明度阶梯', '比色相更影响可读性')
        add_table(s, ['层级', '明度区间', '用途'],
                  [[x.get('level', ''), x.get('range', ''), x.get('usage', '')] for x in lum],
                  col_w=[1, 1, 3], font_size=12, row_h=0.4, max_rows=_cap(1.6, 0.4))
        caption(s, data.get('luminance_note', '把成品转灰度后主体与背景仍可分，明度层级才算合格。'))
        footer(s, foot)
    else:
        skipped.append('明度阶梯')

    # 5 资产清单
    ast = data.get('assets') or []
    if ast:
        s = new_page(prs, blank, '资产清单', '按类别归口，规格与图集分组写全')
        rows = [[a.get('category', ''), a.get('count', ''), a.get('priority', ''),
                 a.get('spec', ''), a.get('group', '')] for a in ast]
        add_table(s, ['类别', '数量', '优先级', '规格要点', '图集分组'], rows,
                  col_w=[0.9, 0.6, 0.7, 2.0, 1.2], font_size=11, row_h=0.36,
                  max_rows=_cap(1.6, 0.36))
        caption(s, data.get('assets_note', '同屏同时出现的资产归同一图集，避免首屏加载无关资源。'))
        footer(s, foot)
    else:
        skipped.append('资产清单')

    # 6 资产数量分布（饼图）
    if ast:
        s = new_page(prs, blank, '资产数量分布', '产能主要压在哪个类别')
        add_chart(s, XL_CHART_TYPE.PIE, [a.get('category', '') for a in ast],
                  [('资产数量', [a.get('count') or 0 for a in ast])],
                  '各资产类别数量占比', pct_labels=True, legend=True)
        caption(s, data.get('assets_chart_note', '哪一类数量最多，就该最先定它的规格与锚点。'))
        footer(s, foot)

    # 7 产出批次计划（条形图 + 表）
    ba = data.get('batches') or []
    if ba:
        s = new_page(prs, blank, '产出批次计划', '先定调，再批量；第 1 批只出 1–2 张')
        add_chart(s, XL_CHART_TYPE.BAR_CLUSTERED, [b.get('batch', '') for b in ba],
                  [('本批数量', [b.get('count') or 0 for b in ba])],
                  '各批次产出数量', legend=False, h=2.6,
                  accent_series=True)
        rows = [[b.get('batch', ''), b.get('content', ''), b.get('count', ''),
                 b.get('goal', '')] for b in ba]
        add_table(s, ['批次', '内容', '数量', '目的'], rows,
                  y=4.4, col_w=[0.7, 3, 0.6, 1.7], font_size=10.5, row_h=0.3,
                  max_rows=_cap(4.4, 0.3))
        caption(s, data.get('batches_note', '第 1 批是基准图，定稿前不要批量生成后面的批次。'))
        footer(s, foot)
    else:
        skipped.append('产出批次计划')

    # 8 缺口清单
    gp = data.get('gaps') or []
    if gp:
        s = new_page(prs, blank, '缺口清单', '每个 P0 都要写清阻塞什么')
        rows = [[g.get('asset', ''), g.get('category', ''), g.get('priority', ''),
                 g.get('blocks', ''), g.get('status', '')] for g in gp]
        add_table(s, ['资产名', '类别', '优先级', '阻塞什么', '状态'], rows,
                  col_w=[1.4, 0.8, 0.6, 1.6, 0.8], font_size=11, row_h=0.36,
                  max_rows=_cap(1.6, 0.36))
        caption(s, data.get('gaps_note', '优先级只按玩法依赖定，不按好不好看定。'))
        footer(s, foot)
    else:
        skipped.append('缺口清单')

    # 9 评审结论与需精修项
    rv = data.get('review') or {}
    if rv:
        s = new_page(prs, blank, '评审结论与需精修项', '结论分档：通过 / 小改 / 重做')
        tf = K.txbox(s, 0.7, 1.5, SW - 1.4, 0.9)
        para(tf, '总体结论：%s' % rv.get('verdict', '—'), size=15, bold=True,
             color=C_ACCENT, first=True, space_after=3)
        if rv.get('summary'):
            para(tf, rv['summary'], size=12.5, color=C_TEXT, space_after=0)
        issues = rv.get('issues') or []
        if issues:
            rows = [[x.get('asset', ''), x.get('issue', ''), x.get('fix', ''),
                     x.get('priority', '')] for x in issues]
            add_table(s, ['资产', '问题', '具体改法', '优先级'], rows,
                      y=2.6, col_w=[1.2, 1.6, 2.2, 0.6], font_size=10.5, row_h=0.34,
                      max_rows=_cap(2.6, 0.34))
        caption(s, '写不出具体改法的，不算问题；涉及风格的判断必须引用风格锚点。')
        footer(s, foot)

    # 10 下一步与待定
    nx = data.get('next') or []
    s = new_page(prs, blank, '下一步与待定事项', '这些定下来之前不要开工')
    tf = K.txbox(s, 0.7, 1.6, SW - 1.4, 4.0)
    if nx:
        for i, t in enumerate(nx):
            para(tf, '• %s' % t, size=13, color=C_TEXT, first=(i == 0), space_after=8)
    else:
        para(tf, '本次无待定事项。', size=14, first=True)
    if skipped:
        para(tf, '⚠️ 以下板块因内容缺失未进入 PPT，需补齐后再出：%s' % '；'.join(skipped),
             size=11.5, color=K.C_WARN, space_after=0)
    footer(s, foot)

    return prs, skipped


def main():
    run_cli(build, _default_out_name, __doc__)


if __name__ == '__main__':
    main()
